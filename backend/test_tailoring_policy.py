import inspect
from io import BytesIO
from unittest.mock import patch

import main
import pytest
from services import scoring_service
from services.tailoring_policy import (
    classify_skill,
    get_tailoring_mode,
    load_adjacency,
    tailoring_prompt,
)

from docx import Document


@pytest.fixture
def resume_bytes():
    doc = Document()
    doc.add_paragraph("Terraform")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Infrastructure as Code"
    table.cell(0, 1).text = "Terraform"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize(
    ("mode", "skill", "truth", "expected"),
    [
        ("strict", "Terraform", {"terraform": "yes"}, "base"),
        ("strict", "Pulumi", {"pulumi": "familiar"}, "blocked"),
        ("balanced", "Pulumi", {"pulumi": "familiar"}, "familiar"),
        ("aggressive", "Pulumi", {"pulumi": "familiar"}, "familiar"),
        ("balanced", "OpenTofu", {}, "adjacent"),
        ("balanced", "UnknownTool", {}, "blocked"),
        ("aggressive", "Pulumi", {"pulumi": "no"}, "blocked"),
    ],
)
def test_mode_skill_matrix(mode, skill, truth, expected):
    disposition, _ = classify_skill(skill, "Delivered Terraform modules", mode, truth)
    assert disposition == expected


def test_nontransferable_language_category_never_grants_adjacency():
    disposition, _ = classify_skill("TypeScript", "Built Python automation", "balanced", {})
    assert disposition == "blocked"


def test_adjacency_loader_is_valid_and_cached():
    load_adjacency.cache_clear()
    first = load_adjacency()
    second = load_adjacency()
    assert first is second
    assert len(first["categories"]) >= 15
    assert any(c["name"] == "Programming languages" and not c["transferable"] for c in first["categories"])


def test_candidate_truth_levels_and_legacy_question_mark(tmp_path, monkeypatch):
    truth_path = tmp_path / "candidate_truth.md"
    truth_path.write_text("- Terraform: yes\n- Pulumi: familiar\n- COBOL: no\n- Kafka: ?\n", encoding="utf-8")
    monkeypatch.setattr(scoring_service, "_TRUTH_PATH", str(truth_path))
    monkeypatch.setattr(scoring_service, "_truth_cache", None)
    monkeypatch.setattr(scoring_service, "_truth_mtime", None)
    assert scoring_service.load_candidate_truth() == {
        "terraform": "yes", "pulumi": "familiar", "cobol": "no", "kafka": "no"
    }


def test_skills_cap_and_no_violation(resume_bytes, monkeypatch):
    monkeypatch.setenv("SKILLS_ADD_CAP", "2")
    truth = {"pulumi": "familiar", "opentofu": "no", "spacelift": "familiar"}
    result = {"skills_additions": [{"category": "Infrastructure as Code", "skills": ["Pulumi", "OpenTofu", "Spacelift"]}]}
    with patch("main.load_candidate_truth", return_value=truth):
        _, added, _, dropped = main._apply_skills_additions(
            resume_bytes, result, base_resume_text="Terraform", tailoring_mode="balanced"
        )
    assert "OpenTofu" not in [x["skill"] for x in added]
    assert len(added) <= 2
    assert any(x["skill"] == "OpenTofu" and x["reason"] == "no" for x in dropped)


def test_aggressive_familiar_is_summary_only_and_once():
    resume = "PROFESSIONAL SUMMARY\nCloud engineer.\nEXPERIENCE\nAcme\nBuilt Python automation."
    reps = [
        {"original": "Cloud engineer.", "new": "Cloud engineer with hands-on exposure to Pulumi.", "keywords_added": ["Pulumi"]},
        {"original": "Built Python automation.", "new": "Built Pulumi automation.", "keywords_added": ["Pulumi"]},
    ]
    with patch("main.load_candidate_truth", return_value={"pulumi": "familiar"}):
        kept, rejected = main._drop_unsupported_tool_claims(reps, ["Pulumi"], resume, "aggressive")
    assert kept == [reps[0]]
    assert rejected[0]["unsupported"] == ["Pulumi"]


def test_retry_and_telegram_paths_forward_mode():
    scan_source = inspect.getsource(main._scan_resume_core)
    batch_source = inspect.getsource(main._process_single_jd)
    telegram_source = inspect.getsource(main._telegram_process_jd_impl)
    assert scan_source.count("tailoring_mode=tailoring_mode") >= 3
    assert batch_source.count("tailoring_mode=tailoring_mode") >= 3
    assert 'settings.get("tailoring_mode")' in telegram_source


def test_mode_default_and_prompt_rules(monkeypatch):
    monkeypatch.delenv("TAILORING_MODE", raising=False)
    assert get_tailoring_mode() == "balanced"
    prompt = tailoring_prompt("aggressive", {"pulumi": "familiar", "cobol": "no"})
    assert "at most ONE summary bullet" in prompt
    assert "never place it in an employer bullet" in prompt
    assert "COBOL".lower() in prompt.lower()
