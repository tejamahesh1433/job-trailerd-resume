"""
Regression tests for JD skills tailoring. The resume's Technical Proficiency block is an
11-row category | skills table, but the tailoring prompt used to ask the AI to "rewrite the
Skills line" — which doesn't exist as one paragraph — so skills additions depended on the
model reproducing a ~200-char table cell character-for-character, and nothing verified the
JD's missing keywords actually ended up in the resume. Skills are now added by code
(add_skills_to_docx, additive-only) and the result is verified (find_absent_keywords).
"""
import os
from io import BytesIO

import docx
import pytest

from services.docx_service import (
    add_skills_to_docx,
    extract_skills_table,
    extract_text_from_docx,
    find_absent_keywords,
)
import main as main_module

BASE_RESUME = os.path.join(os.path.dirname(__file__), "original", "base_resume.docx")


@pytest.fixture(scope="module")
def resume_bytes():
    with open(BASE_RESUME, "rb") as f:
        return f.read()


def _cell_for(file_bytes, category):
    table = docx.Document(BytesIO(file_bytes)).tables[0]
    row = next(r for r in table.rows if r.cells[0].text.strip() == category)
    return row.cells[1].text


def test_extract_skills_table_returns_real_category_names(resume_bytes):
    table = extract_skills_table(resume_bytes)
    categories = [r["category"] for r in table]
    assert len(table) == 11
    assert "Monitoring & Observability" in categories
    assert "Cloud Platforms" in categories
    assert "Terraform" in next(r["skills"] for r in table if r["category"] == "Infrastructure as Code")


def test_new_skill_is_appended_to_its_category_and_existing_skills_are_untouched(resume_bytes):
    before = _cell_for(resume_bytes, "Monitoring & Observability")
    stream, added, unplaced = add_skills_to_docx(
        resume_bytes, [{"category": "Monitoring & Observability", "skills": ["New Relic", "Thanos"]}])
    after = _cell_for(stream.read(), "Monitoring & Observability")

    assert added == [("Monitoring & Observability", "New Relic"), ("Monitoring & Observability", "Thanos")]
    assert unplaced == []
    assert after == before + ", New Relic, Thanos"


def test_skills_already_in_the_table_are_not_duplicated(resume_bytes):
    # Terraform + Ansible already listed under IaC; case differences must not defeat the check
    stream, added, _ = add_skills_to_docx(
        resume_bytes, [{"category": "Infrastructure as Code", "skills": ["terraform", "ANSIBLE", "Pulumi"]}])
    assert [s for _, s in added] == ["Pulumi"]
    assert _cell_for(stream.read(), "Infrastructure as Code").lower().count("terraform") == 1


def test_a_skill_listed_under_a_different_category_is_still_not_duplicated(resume_bytes):
    # Docker is already in Containers & Orchestration; the AI mis-filing it under CI/CD must not add a 2nd copy
    _, added, _ = add_skills_to_docx(resume_bytes, [{"category": "CI/CD & DevOps Tools", "skills": ["Docker"]}])
    assert added == []


def test_applying_the_same_additions_twice_is_idempotent(resume_bytes):
    additions = [{"category": "Networking", "skills": ["Cloudflare"]}]
    once = add_skills_to_docx(resume_bytes, additions)[0].read()
    stream, added, _ = add_skills_to_docx(once, additions)
    assert added == []
    assert _cell_for(stream.read(), "Networking").count("Cloudflare") == 1


def test_category_name_is_fuzzy_matched(resume_bytes):
    _, added, unplaced = add_skills_to_docx(
        resume_bytes, [{"category": "Monitoring and Observability", "skills": ["Thanos"]}])
    assert added == [("Monitoring & Observability", "Thanos")]
    assert unplaced == []


def test_unknown_category_is_reported_not_silently_dropped(resume_bytes):
    stream, added, unplaced = add_skills_to_docx(
        resume_bytes, [{"category": "Quantum Basket Weaving", "skills": ["Qiskit"]}])
    assert added == []
    assert unplaced == ["Qiskit"]
    assert "Qiskit" not in extract_text_from_docx(stream.read())


def test_new_text_inherits_the_cell_run_formatting(resume_bytes):
    def last_run(file_bytes):
        table = docx.Document(BytesIO(file_bytes)).tables[0]
        row = next(r for r in table.rows if r.cells[0].text.strip() == "Networking")
        return next(r for r in reversed(row.cells[1].paragraphs[-1].runs) if r.text)

    original_run = last_run(resume_bytes)
    modified = add_skills_to_docx(resume_bytes, [{"category": "Networking", "skills": ["Cloudflare"]}])[0].read()
    new_run = last_run(modified)
    assert "Cloudflare" in new_run.text
    assert (new_run.font.name, new_run.font.size, new_run.bold) == (original_run.font.name, original_run.font.size, original_run.bold)


def test_resume_without_a_skills_table_reports_everything_as_unplaced():
    d = docx.Document()
    d.add_paragraph("Some Person")
    buf = BytesIO()
    d.save(buf)
    assert extract_skills_table(buf.getvalue()) == []
    _, added, unplaced = add_skills_to_docx(buf.getvalue(), [{"category": "Networking", "skills": ["Cloudflare"]}])
    assert added == [] and unplaced == ["Cloudflare"]


def test_find_absent_keywords_matches_whole_terms_only(resume_bytes):
    absent = find_absent_keywords(resume_bytes, ["Kubernetes", "Kafka", "Go", "kubernetes", "  ", "Kafka"])
    # Kubernetes present (any case); Kafka absent (reported once); "Go" must NOT match "Google"/"GoLang" fragments
    assert absent == ["Kafka", "Go"]


def test_clean_skills_additions_drops_malformed_ai_output():
    raw = [
        {"category": "Networking", "skills": ["Cloudflare", "", 5, None]},
        {"category": "Networking"},
        {"skills": ["x"]},
        "nonsense",
        {"category": 7, "skills": ["y"]},
    ]
    assert main_module._clean_skills_additions(raw) == [{"category": "Networking", "skills": ["Cloudflare"]}]
    assert main_module._clean_skills_additions(None) == []
    assert main_module._clean_skills_additions({"category": "x"}) == []


def test_apply_skills_additions_end_to_end_reports_added_and_unmatched(resume_bytes):
    result = {
        "skills_additions": [{"category": "Monitoring & Observability", "skills": ["New Relic"]}],
        # New Relic gets placed; Angular is the non-DevOps keyword the prompt forbids adding, Kafka the AI never
        # placed; Java is already in the resume's bullets so it must NOT be flagged as absent.
        "missing_keywords": ["New Relic", "Angular", "Kafka", "Java"],
    }
    new_bytes, skills_added, unmatched = main_module._apply_skills_additions(resume_bytes, result)

    assert skills_added == [{"category": "Monitoring & Observability", "skill": "New Relic"}]
    assert unmatched == ["Angular", "Kafka"]
    assert "New Relic" in extract_text_from_docx(new_bytes)


def test_apply_skills_additions_with_no_ai_additions_still_verifies_keywords(resume_bytes):
    new_bytes, skills_added, unmatched = main_module._apply_skills_additions(
        resume_bytes, {"missing_keywords": ["Kafka", "Terraform"]})
    assert new_bytes == resume_bytes
    assert skills_added == []
    assert unmatched == ["Kafka"]


@pytest.mark.parametrize("phrase", [
    "11+ years", "5 years of DevOps", "financial services domain", "banking experience",
    "Banking-Approved Compliance Zones and Regions", "Security clearance required", "Bachelor degree",
])
def test_requirement_phrases_are_not_treated_as_skills(phrase):
    assert not main_module._looks_like_skill(phrase)


@pytest.mark.parametrize("skill", [
    "Puppet", "Kafka", "GitHub Actions", "Blue-Green Deployments", "CI/CD", "Node.js", "C#", "Azure DevOps",
])
def test_real_tools_pass_the_skill_filter(skill):
    assert main_module._looks_like_skill(skill)


def test_clean_skills_additions_drops_requirement_phrases_but_keeps_real_skills():
    raw = [{"category": "Cloud Platforms", "skills": ["Puppet", "11+ years", "fintech domain"]},
           {"category": "Networking", "skills": ["financial services domain"]}]
    assert main_module._clean_skills_additions(raw) == [{"category": "Cloud Platforms", "skills": ["Puppet"]}]


# ── TOOL-CLAIM RULE: bullets may not claim tools the base resume never mentions ──────────────
RESUME_FOR_CLAIMS = (
    "Senior DevOps Engineer. 10+ years of DevOps and SRE experience across AWS, Azure and GCP. "
    "Automated configuration management using Ansible Tower and Chef. Built CI/CD with Jenkins on Google Cloud. "
    "Infrastructure as Code (IaC) with Terraform. Financial Group, banking platforms."
)


def _claims(rep, jd_keywords=()):
    return main_module._unsupported_tool_keywords(rep, list(jd_keywords), RESUME_FOR_CLAIMS)


def test_bullet_naming_a_tool_absent_from_the_resume_is_flagged():
    rep = {"new": "Automated configuration management using Ansible Tower, Chef, and Puppet.",
           "keywords_added": ["Puppet"]}
    assert _claims(rep) == ["Puppet"]


def test_tool_is_flagged_even_if_the_ai_did_not_self_report_it_but_it_is_a_missing_keyword():
    rep = {"new": "Automated configuration management using Ansible Tower, Chef, and Puppet.",
           "keywords_added": []}
    assert _claims(rep, jd_keywords=["Puppet", "Kafka"]) == ["Puppet"]  # Kafka isn't in the new text


def test_rewording_around_tools_already_in_the_resume_is_allowed():
    rep = {"new": "Automated IaC at scale using Terraform, Ansible Tower and Chef with CI/CD on Jenkins.",
           "keywords_added": ["IaC", "Terraform", "CI/CD", "Ansible"]}
    assert _claims(rep) == []


def test_compound_terms_are_judged_word_by_word():
    ok = {"new": "10+ years of DevOps/SRE experience.", "keywords_added": ["DevOps/SRE"]}
    bad = {"new": "Enforced compliance zones.", "keywords_added": ["compliance zones"]}
    assert _claims(ok) == []
    assert _claims(bad) == ["compliance zones"]


def test_years_and_industry_wording_are_not_tool_claims():
    rep = {"new": "11+ years delivering fintech and banking platforms.",
           "keywords_added": ["11+ years", "fintech", "banking", "financial services"]}
    assert _claims(rep) == []


def test_short_tool_name_does_not_match_inside_a_longer_word():
    # resume has "Google" but not the Go language — "Go" must not count as already present
    rep = {"new": "Built internal tooling in Go.", "keywords_added": ["Go"]}
    assert _claims(rep) == ["Go"]


def test_drop_unsupported_tool_claims_keeps_good_edits_and_reports_the_rest():
    good = {"original": "a", "new": "Led Terraform IaC rollouts.", "keywords_added": ["Terraform", "IaC"]}
    bad = {"original": "b", "new": "Ran Chef and Puppet.", "keywords_added": ["Puppet"]}
    kept, rejected = main_module._drop_unsupported_tool_claims([good, bad], ["Puppet"], RESUME_FOR_CLAIMS)
    assert kept == [good]
    assert [(r["original"], r["unsupported"]) for r in rejected] == [("b", ["Puppet"])]
